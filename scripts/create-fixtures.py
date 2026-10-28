"""Synthetic visual fixtures; no user documents."""
from pathlib import Path
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
from openpyxl import Workbook
root=Path('.agent-state/fixtures');root.mkdir(parents=True,exist_ok=True)
font=ImageFont.truetype(next(str(p) for p in [Path('/System/Library/Fonts/Supplemental/Arial.ttf'),Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),Path('/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf')] if p.exists()),36)
im=Image.new('RGB',(1000,800),'white');d=ImageDraw.Draw(im);d.rectangle((0,0,499,399),fill='red');d.rectangle((500,0,999,399),fill='blue');d.rectangle((0,400,499,799),fill='green');d.text((540,540),'Visual fixture',fill='black',font=font);im.save(root/'quadrants.png')
pages=[]
for i in range(2):
 p=Image.new('RGB',(1000,1400),'white');q=ImageDraw.Draw(p);q.text((100,100),'SYNTHETIC PAGE '+str(i+1),fill='black',font=font);q.rectangle((100,250,850,650),fill=('royalblue' if i==0 else 'tomato'));pages.append(p)
pages[0].save(root/'two-pages.pdf',save_all=True,append_images=pages[1:],resolution=100)
doc=Document();doc.add_heading('视觉复查 / Visual Review',0);doc.add_paragraph('Synthetic document; no private data. 检查中文、表格和分页。');table=doc.add_table(rows=3,cols=2);table.cell(0,0).text='Item';table.cell(0,1).text='Status';table.cell(1,0).text='Background renderer';table.cell(1,1).text='Fixture';doc.add_page_break();doc.add_heading('Second page / 第二页',1);doc.add_paragraph('Only selected page numbers should be returned.');doc.save(root/'sample.docx')
prs=Presentation()
for i in range(2):
 slide=prs.slides.add_slide(prs.slide_layouts[5]);slide.shapes.title.text='视觉复查 · Slide '+str(i+1);shape=slide.shapes.add_shape(1, Inches(1), Inches(2), Inches(5), Inches(2));shape.fill.solid();shape.fill.fore_color.rgb=RGBColor(30,100,190);shape.text='SYNTHETIC VISUAL EVIDENCE'
prs.save(root/'sample.pptx')
w=Workbook();s=w.active;s.title='检查表';s.append(['项目 / Item','状态 / Status']);s.append(['Renderer','Synthetic']);s.append(['Pages','Print layout']);s.column_dimensions['A'].width=25;s.column_dimensions['B'].width=30;s.print_area='A1:B3';s.page_setup.orientation='landscape';w.save(root/'sample.xlsx')
(root/'sample.md').write_text('# 视觉复查 / Markdown\n\n|Item|Status|\n|---|---|\n|Rendering|Synthetic|\n')
(root/'sample.html').write_text('<!doctype html><meta charset="utf-8"><h1>后台预览 / HTML</h1><div style="width:250px;height:100px;background:#3366aa;color:white">SYNTHETIC BLOCK</div><script>document.body.textContent="SHOULD NEVER RUN"</script><img src="http://host.docker.internal:18520/should-not-load">')
print('Created seven synthetic file types')
